from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook

from app.ingestion.base import DocumentParser
from app.ingestion.models import Document, DocumentElement


class XLSXParser(DocumentParser):

    def parse(self, file_path: Path) -> Document:
        document_id = str(uuid4())

        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )

        elements = []

        for worksheet in workbook.worksheets:

            for row_number, row in enumerate(
                worksheet.iter_rows(values_only=True),
                start=1,
            ):
                values = [
                    str(value).strip()
                    for value in row
                    if value is not None
                ]

                if not values:
                    continue

                text = " | ".join(values)

                elements.append(
                    DocumentElement(
                        element_id=str(uuid4()),
                        element_type="row",
                        text=text,
                        metadata={
                            "sheet": worksheet.title,
                            "row": row_number,
                        },
                    )
                )

        workbook.close()

        return Document(
            document_id=document_id,
            file_name=file_path.name,
            file_type="xlsx",
            elements=elements,
        )