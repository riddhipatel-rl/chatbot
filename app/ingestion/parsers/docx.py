from pathlib import Path
from uuid import uuid4

from docx import Document as DocxDocument

from app.ingestion.base import DocumentParser
from app.ingestion.models import Document, DocumentElement


class DOCXParser(DocumentParser):

    def parse(self, file_path: Path) -> Document:
        document_id = str(uuid4())
        docx = DocxDocument(file_path)

        elements = []

        for index, paragraph in enumerate(docx.paragraphs):
            text = paragraph.text.strip()

            if not text:
                continue

            element_type = "heading" if paragraph.style.name.startswith("Heading") else "text"

            elements.append(
                DocumentElement(
                    element_id=str(uuid4()),
                    element_type=element_type,
                    text=text,
                    metadata={
                        "element_index": index,
                    },
                )
            )

        return Document(
            document_id=document_id,
            file_name=file_path.name,
            file_type="docx",
            elements=elements,
        )