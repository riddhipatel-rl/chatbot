from pathlib import Path
from uuid import uuid4

from pptx import Presentation

from app.ingestion.base import DocumentParser
from app.ingestion.models import Document, DocumentElement


class PPTXParser(DocumentParser):

    def parse(self, file_path: Path) -> Document:
        document_id = str(uuid4())
        presentation = Presentation(file_path)

        elements = []

        for slide_number, slide in enumerate(presentation.slides, start=1):

            for shape_index, shape in enumerate(slide.shapes):

                if not hasattr(shape, "text"):
                    continue

                text = shape.text.strip()

                if not text:
                    continue

                elements.append(
                    DocumentElement(
                        element_id=str(uuid4()),
                        element_type="text",
                        text=text,
                        metadata={
                            "slide": slide_number,
                            "shape_index": shape_index,
                        },
                    )
                )

        return Document(
            document_id=document_id,
            file_name=file_path.name,
            file_type="pptx",
            elements=elements,
        )